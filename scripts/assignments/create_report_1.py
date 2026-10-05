"""Build a template-preserving report with selectable prompts and full JSON."""
from pathlib import Path
from copy import deepcopy
from io import BytesIO
import hashlib
import json
import zipfile
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[2]
SOURCE = ROOT / '课程文件/作业报告模版.docx'
OUT = ROOT / 'docs/assignments/提交/524031910745-作业1.docx'
assert hashlib.sha256(SOURCE.read_bytes()).hexdigest() == 'ca56585e27ae3aa39fc50657b3fc0c48027f8cee49aceea2de5ddde4b2585e23'
doc = Document(SOURCE)
heading = deepcopy(doc.paragraphs[4]._p)
doc.paragraphs[0].text = '2026-10-06'
for run in doc.paragraphs[1].runs:
    run.text = run.text.replace('X', '1')
for paragraph in list(doc.paragraphs)[4:]:
    paragraph._p.getparent().remove(paragraph._p)

def prose(text, bold=False, code=False, keep=False):
    p = doc.add_paragraph(style='Normal')
    p.paragraph_format.keep_with_next = keep
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(6 if not code else 0)
    if code:
        p.paragraph_format.left_indent = Pt(0)
        p.paragraph_format.line_spacing = Pt(11)
    run = p.add_run(text)
    run.font.color.rgb = RGBColor(0, 0, 0)
    run.bold = bold
    if code:
        run.font.name = 'Courier New'
        run.font.size = Pt(9)
    return p

def section(text):
    p = deepcopy(heading)
    for child in list(p):
        if child.tag != qn('w:pPr'):
            p.remove(child)
    props = p.find(qn('w:pPr'))
    props.append(OxmlElement('w:keepNext'))
    r = OxmlElement('w:r')
    rp = OxmlElement('w:rPr')
    color = OxmlElement('w:color'); color.set(qn('w:val'), '000000'); rp.append(color)
    r.append(rp); t = OxmlElement('w:t'); t.text = text; r.append(t); p.append(r)
    doc._element.body.insert(len(doc._element.body)-1, p)

def compact(value, depth=0):
    short = json.dumps(value, ensure_ascii=False, separators=(',', ': '))
    pad = '  ' * depth
    if len(short) + len(pad) <= 88 or not isinstance(value, (dict, list)):
        return pad + short
    if isinstance(value, dict):
        rows = []
        for key, item in value.items():
            rendered = compact(item, depth+1).splitlines()
            rendered[0] = '  '*(depth+1) + json.dumps(key,ensure_ascii=False) + ': ' + rendered[0].lstrip()
            rows.append('\n'.join(rendered))
        return pad+'{\n'+',\n'.join(rows)+'\n'+pad+'}'
    return pad+'[\n'+',\n'.join(compact(item,depth+1) for item in value)+'\n'+pad+']'

directory = ROOT / 'docs/assignments/作业1'
transcript = json.loads((directory/'conversation.json').read_text())
evidence = json.loads((ROOT/'docs/assignments/evidence/1e-contract.json').read_text())
canonical_path = ROOT / 'backend/src/main/resources/static/api/v1/openapi.json'
canonical = canonical_path.read_bytes()
section('完整提示词与实现契约')
prose(f"本作业使用真实 DeepSeek API，为知页书城已实现的四个 REST 接口生成 OpenAPI 3.0.3 JSON。角色提示限定架构师职责，任务提示提供实际字段、默认值和错误边界；通过{len(transcript['rounds'])}轮记录与独立校验确认最终规范可用于接口测评。")
for line in ['GET /api/books  获取分页书目与分类过滤', 'GET /api/books/{id}  查看单本详情', 'POST /api/books  添加新书', 'PATCH /api/books/{id}/inventory  仅更新库存']:
    prose(line)
prose('System Prompt', bold=True, keep=True)
for line in transcript['system'].splitlines():
    if line: prose(line)
prose('User Prompt', bold=True, keep=True)
for line in transcript['rounds'][0]['user'].splitlines():
    if line: prose(line)

section('多轮生成与修订记录')
prose('服务商 DeepSeek，模型 deepseek-flash，调用端点 https://api.deepseek.com/chat/completions。使用非思考模式、temperature=0 与 JSON 输出模式。完整原始消息、响应 ID、用量和时间保存在 conversation.json；记录不含 API 密钥。')
prose('前三轮保存结果通过 OpenAPI 结构检查。响应校验器随后发现错误 data 缺少明确类型，第四轮反馈真实报错并修正；同时修正重复 ID 示例码为40903，并允许历史图书响应 stock 为 null。第五轮根据复核修正创建库存描述和一般参数错误示例码40000。各轮路径、方法及核心请求约束保持一致。低温度参数与 Few-Shot 约束用于减少漂移，不代表模型逐字节确定。')
for round in transcript['rounds'][1:]:
    prose(f"第 {round['round']} 轮完整 User Prompt", bold=True, keep=True)
    prose(round['user'])
for number in [1,2,4]:
    image = ROOT / f'docs/assignments/evidence/1e-conversation-{number}.png'
    assert image.exists(), f'真实截图尚未生成：{image}'
    prose(f'第 {number} 轮真实 API 对话截图', bold=True, keep=True)
    p = doc.add_paragraph(style='Normal')
    p.paragraph_format.left_indent = Pt(0)
    drawing = p.add_run().add_picture(str(image), width=Inches(5.76))
    drawing._inline.docPr.set('descr', f'DeepSeek API 第{number}轮真实输入与响应日志')
    prose('截图由本地原始日志查看器显示，非 DeepSeek 官方网页；为可读性仅添加 JSON 缩进，完整内容见前述提示词及下方最终规范。')

section('前后端测评结果')
prose(f"2026-10-06，本机真实 Spring Boot 与 MySQL 环境下，{len(evidence['http'])} 个 HTTP 响应按最终 schema 校验通过，含分页成功及空页、非法参数400、详情404、新增201与Location、重复ID409、库存更新与严格整数拒绝。验收临时图书已清理；线上真实销售库存不在本作业范围内。")
prose('后端48个隔离测试全部通过；前端生产构建通过。桌面浏览器此前逐项完成分页/分类/搜索/排序、空结果、断网重试、详情404、新增失败保留输入与新增持久化、库存24→0→24及缺货禁用加购。PATCH来源检查曾发现403，修正CORS并重测后才进入本功能。按用户要求不测试手机比例。')
prose('OpenAPI静态地址 /api/v1/openapi.json 与最终文件逐字段一致。未人为制造真实500故障，500是已实现异常处理约定。证据目录为 docs/assignments/evidence；完整测试脚本为 scripts/assignments/verify_openapi.mjs。')
prose('最终规范 SHA-256：'+hashlib.sha256(canonical).hexdigest(), code=True)
prose('参考：OpenAPI 3.0.3规范 https://spec.openapis.org/oas/v3.0.3.html；DeepSeek官方JSON与模型说明 https://api-docs.deepseek.com/guides/json_mode/ 及 https://api-docs.deepseek.com/zh-cn/。')

section('最终 OpenAPI JSON 全文')
prose('以下为最终模型生成并通过真实响应校验的完整 JSON。独立规范文件位于 backend/src/main/resources/static/api/v1/openapi.json，可从后端静态地址下载。')
formatted = compact(json.loads(canonical))
assert json.loads(formatted) == json.loads(canonical)
for line in formatted.splitlines():
    prose(line,code=True)

# python-docx authors the intended body/images. Repack untouched original parts
# byte-for-byte to preserve template styles, numbering, footers and settings.
buffer = BytesIO(); doc.save(buffer)
original = zipfile.ZipFile(SOURCE)
generated = zipfile.ZipFile(BytesIO(buffer.getvalue()))
changed = {'word/document.xml','word/_rels/document.xml.rels','[Content_Types].xml'}
OUT.parent.mkdir(parents=True,exist_ok=True)
with zipfile.ZipFile(OUT,'w',zipfile.ZIP_DEFLATED) as final:
    for name in original.namelist():
        final.writestr(name,generated.read(name) if name in changed else original.read(name))
    for name in generated.namelist():
        if name not in original.namelist(): final.writestr(name,generated.read(name))
with zipfile.ZipFile(OUT) as final:
    for name in original.namelist():
        if name not in changed: assert original.read(name)==final.read(name),name
assert hashlib.sha256(SOURCE.read_bytes()).hexdigest() == 'ca56585e27ae3aa39fc50657b3fc0c48027f8cee49aceea2de5ddde4b2585e23'
print(f'报告已创建，JSON可选文本{len(formatted.splitlines())}行，模板保留部件逐字节验证通过：{OUT}')
